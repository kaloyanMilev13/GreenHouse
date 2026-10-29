from pathlib import Path
ns={}
exec(Path('tmp/pdfs/build_plan.py').read_text().split('flow=[]')[0],ns)
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from html.parser import HTMLParser
D=Document(); sec=D.sections[0]
sec.page_width=Pt(595.28);sec.page_height=Pt(841.89)
sec.top_margin=Pt(30);sec.bottom_margin=Pt(32);sec.left_margin=Pt(40);sec.right_margin=Pt(40)
for name in ['Normal','Title','Heading 1']:
 s=D.styles[name];s.font.name='Noto Sans';s.font.color.rgb=RGBColor(0,0,0)
s=D.styles['Normal'];s.font.size=Pt(9.5);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.0
s=D.styles['Title'];s.font.size=Pt(22);s.font.bold=True;s.paragraph_format.space_after=Pt(8)
s=D.styles['Heading 1'];s.font.size=Pt(14);s.font.bold=True;s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(7)
class H(HTMLParser):
 def __init__(self,p,size=None):super().__init__();self.p=p;self.bold=False;self.url=None;self.size=size
 def handle_starttag(self,t,a):
  if t=='b':self.bold=True
  if t=='a':self.url=dict(a)['href']
  if t=='br':self.p.add_run().add_break()
 def handle_endtag(self,t):
  if t=='b':self.bold=False
  if t=='a':self.url=None
 def handle_data(self,text):
  r=self.p.add_run(text);r.bold=self.bold
  if self.size:r.font.size=Pt(self.size)
  if self.url:
   r.font.color.rgb=RGBColor.from_string('146E80');r.font.underline=True
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),self.p.part.relate_to(self.url,RT.HYPERLINK,is_external=True));h.append(r._r);self.p._p.append(h)
def para(t,style=None):
 p=D.add_paragraph(style=style);H(p).feed(t);return p
def table(rows):
 t=D.add_table(rows=1,cols=5);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[45,111,125,113,121]
 for c,w in zip(t.columns,widths):c.width=Pt(w)
 heads=['Period','Learning objectives','Resource and reason for choosing it','Practical application','Deliverable / completion criteria']
 for i,values in enumerate([heads]+rows):
  row=t.rows[0] if i==0 else t.add_row()
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,(cell,txt) in enumerate(zip(row.cells,values)):
   cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'24485A' if i==0 else ('EEF3F6' if i%2 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
   H(p,8.5).feed(txt.split('<br/>')[0] if j==0 else txt)
   if i==0:
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
para('Embedded Systems study plan','Title');para('FONTYS ICT / FIRST YEAR / COACH: IMAN')
para('A. Rationale - Why Embedded Systems?','Heading 1');para(ns['rationale'])
para('B. Learning roadmap - 12 weeks','Heading 1');para(ns['intro']);table(ns['rows'][:3])
D.add_page_break();para('B. Learning roadmap - continued','Heading 1')
para('Order: quantities and circuits > C and memory > timing > network reliability > calculus > feedback. Each phase receives 9 hours. Later work builds on the relevant completion gates.')
table(ns['rows'][3:])
D.add_page_break();para('C. Weekly routine','Heading 1')
for name,txt in ns['weekly']:para('<b>'+name+'</b> - '+txt)
para('Total: 4.5 h/week. A fifth hour is optional catch-up, not new content. If a gate fails, use the next practice session to repair the gap and reduce later scope. If necessary, finish week 12 with P control and defer PI; do not compress the foundations.')
para('D. Longer-term direction','Heading 1')
for t in ns['longterm']:para(t)
para('E. Progress evaluation','Heading 1')
for i,t in enumerate(ns['evalitems'],1):para(str(i)+'. '+t)
para('Resource access and prerequisites','Heading 1')
para('Resource titles, linked content and relevance checked on 29 October 2026. The core readings/video pages are publicly accessible; no paid course enrolment or certificate is required. Study only the named sections. Entry requirements in the table are my study readiness checks, not provider-issued admissions requirements. GNU material includes extensions: focus on portable C and verify target-specific details. Espressif API versions vary; select the installed framework/chip version. The ROS2 rendered documentation blocked automated access, so its official source was verified instead.')
Path('output/docx').mkdir(exist_ok=True)
D.save('output/docx/embedded-systems-study-plan.docx')
