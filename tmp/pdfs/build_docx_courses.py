from pathlib import Path
ns={}
exec(Path('tmp/pdfs/build_plan.py').read_text().split('flow=[]')[0],ns)
exec(Path('tmp/pdfs/course_revision.py').read_text(),globals())
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from html.parser import HTMLParser
D=Document(); sec=D.sections[0]
sec.page_width=Pt(595.28);sec.page_height=Pt(841.89)
sec.top_margin=Pt(30);sec.bottom_margin=Pt(32);sec.left_margin=Pt(40);sec.right_margin=Pt(40)
for name in ['Normal','Title','Heading 1']:
 s=D.styles[name];s.font.name='Noto Sans';s.font.color.rgb=RGBColor(0,0,0)
s=D.styles['Normal'];s.font.size=Pt(9.5);s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.0
s=D.styles['Title'];s.font.size=Pt(22);s.font.bold=True;s.paragraph_format.space_after=Pt(8)
s=D.styles['Heading 1'];s.font.size=Pt(14);s.font.bold=True;s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(7)
class H(HTMLParser):
 def __init__(self,p,size=None):super().__init__();self.p=p;self.bold=False;self.url=None;self.size=size
 def handle_starttag(self,t,a):
  if t=='b':self.bold=True
  if t=='a':self.url=dict(a)['href']
  if t=='br':self.p.add_run().add_break()
 def handle_endtag(self,t):
  if t=='b':self.bold=False
  if t=='a':self.url=None
 def handle_data(self,text):
  r=self.p.add_run(text);r.bold=self.bold
  if self.size:r.font.size=Pt(self.size)
  if self.url:
   r.font.color.rgb=RGBColor.from_string('146E80');r.font.underline=True
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),self.p.part.relate_to(self.url,RT.HYPERLINK,is_external=True));h.append(r._r);self.p._p.append(h)
def para(t,style=None):
 p=D.add_paragraph(style=style);H(p).feed(t);return p
def table(rows):
 t=D.add_table(rows=1,cols=5);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[45,111,125,113,121]
 for c,w in zip(t.columns,widths):c.width=Pt(w)
 heads=['Period','Learning objectives','Resource and reason for choosing it','Practical application','Deliverable / completion criteria']
 for i,values in enumerate([heads]+rows):
  row=t.rows[0] if i==0 else t.add_row()
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,(cell,txt) in enumerate(zip(row.cells,values)):
   cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'24485A' if i==0 else ('EEF3F6' if i%2 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');mar.append(el)
   pr.append(mar)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
   H(p,8.5).feed(txt.split('<br/>')[0] if j==0 else txt)
   if i==0:
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
para('Embedded Systems study plan','Title');para('FONTYS ICT / FIRST YEAR / COACH: IMAN')
para('A. Rationale - Why Embedded Systems?','Heading 1');para(ns['rationale'])
para('B. Learning roadmap - 12 weeks','Heading 1');para(ns['intro']);table(ns['rows'][:3])
D.add_page_break();para('B. Learning roadmap - continued','Heading 1')
para('Sequence: IoT foundations > C readiness > build tools and memory > independent application. Each fortnight receives 9 hours; course assignments replace separate exercises where they overlap.')
table(ns['rows'][3:])
D.add_page_break();para('C. Weekly routine','Heading 1')
for name,txt in ns['weekly']:para('<b>'+name+'</b> - '+txt)
para('The optional fifth hour is for catch-up. Protect the two-course minimum: reduce the final experiment if assignments take longer. Complete both courses before starting a third. PID implementation and ROS2 move beyond week 12; first develop functions, derivatives, integrals and sampling concepts.')
para('D. Longer-term direction','Heading 1')
for t in ns['longterm']:para(t)
para('E. Progress evaluation','Heading 1')
for i,t in enumerate(ns['evalitems'],1):para(str(i)+'. '+t)
para('Course workload and access','Heading 1')
para('Verified on 29 October 2026. Course 1 requires basic computer knowledge: its headline estimate is 10 hours, while the four module estimates total 12; I budget 14. Course 2 is intermediate and lists 9 hours; I budget 20 for setup, C revision and assignments. Its first course requires no hardware purchase; a Linux/VM development environment is part of the work. These are estimates, not guaranteed completion times.')
para('Both courses are listed for enrolment. Full assessment/certificate access may require payment; check the available enrolment or financial-aid options before starting. This plan assumes access to required assessments and does not assume free certificates. Allow for peer grading. An optional third course belongs after the two-course minimum and the foundational gaps have been reviewed.')
Path('output/docx').mkdir(exist_ok=True)
for root in [D._element,D.styles.element]:
 for el in list(root.iter()):
  if el.tag.endswith('}pBdr'):el.getparent().remove(el)
D.save('output/docx/embedded-systems-study-plan.docx')
